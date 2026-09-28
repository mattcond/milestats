import csv

from openpyxl import load_workbook

from milestats.main import (
    aggrega_per_bucket_temporale,
    calcola_velocita,
    esporta,
    etichetta_bucket_discreti,
    etichetta_categoria_passo,
    filtra_per_finestre,
    righe_export,
)


def _bucket_list(punti_sintetici, mode='a'):
    punti = filtra_per_finestre(list(punti_sintetici), None)
    punti = calcola_velocita(punti, mode)
    bucket_list = aggrega_per_bucket_temporale(punti, 5, mode)
    bucket_list = etichetta_bucket_discreti(bucket_list, 5, 500)
    return etichetta_categoria_passo(bucket_list, mode)


def test_righe_export_intestazione_modalita_a(punti_sintetici):
    bucket_list = _bucket_list(punti_sintetici, mode='a')
    intestazione, righe = righe_export(bucket_list, 'a', 'id123')
    assert intestazione[:2] == ['id_attivita', 'timestamp']
    assert 'time_bucket_idx' in intestazione
    assert 'distance_bucket_idx' in intestazione
    assert 'pace_categoria' in intestazione
    for sigla in ('dev', 'hav', 'uff'):
        assert f'speed_{sigla}_kmh' in intestazione
        assert f'pace_{sigla}_min_km' in intestazione
    assert len(righe) == len(bucket_list)
    assert all(riga[0] == 'id123' for riga in righe)


def test_righe_export_modalita_s_solo_colonne_uff(punti_sintetici):
    bucket_list = _bucket_list(punti_sintetici, mode='s')
    intestazione, _ = righe_export(bucket_list, 's', 'id123')
    assert 'speed_uff_kmh' in intestazione
    assert 'speed_dev_kmh' not in intestazione
    assert 'speed_hav_kmh' not in intestazione


def test_esporta_csv(punti_sintetici, tmp_path):
    bucket_list = _bucket_list(punti_sintetici, mode='s')
    intestazione, righe = righe_export(bucket_list, 's', 'id123')
    percorso = esporta(intestazione, righe, 'csv', tmp_path, 'nome_test')

    assert percorso.exists()
    with open(percorso, newline='', encoding='utf-8') as fh:
        lette = list(csv.reader(fh))
    assert lette[0] == intestazione
    assert len(lette) == 1 + len(righe)


def test_esporta_xlsx(punti_sintetici, tmp_path):
    bucket_list = _bucket_list(punti_sintetici, mode='a')
    intestazione, righe = righe_export(bucket_list, 'a', 'id123')
    percorso = esporta(intestazione, righe, 'xlsx', tmp_path, 'nome_test')

    assert percorso.exists()
    wb = load_workbook(percorso)
    ws = wb.active
    assert [c.value for c in ws[1]] == intestazione
    assert ws.max_row == 1 + len(righe)
    assert ws['A1'].font.bold is True
    assert ws.freeze_panes == "A2"
    idx_timestamp = intestazione.index('timestamp') + 1
    assert ws.cell(row=2, column=idx_timestamp).number_format == 'yyyy-mm-dd hh:mm:ss'
