import csv
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from milestats.main import (
    _deserializza_valore_csv,
    aggiungi_colonna_ultima_attivita,
    leggi_righe_export,
    main,
    unisci_output,
    unisci_righe_export,
)


def _scrivi_csv(percorso, intestazione, righe):
    with open(percorso, 'w', newline='', encoding='utf-8') as fh:
        w = csv.writer(fh)
        w.writerow(intestazione)
        w.writerows(righe)


def _scrivi_xlsx(percorso, intestazione, righe):
    wb = Workbook()
    ws = wb.active
    ws.append(intestazione)
    for riga in righe:
        ws.append(riga)
    wb.save(percorso)


def test_deserializza_valore_csv_ricostruisce_i_tipi():
    assert _deserializza_valore_csv('') is None
    assert _deserializza_valore_csv('42') == 42
    assert _deserializza_valore_csv('6.302') == 6.302
    assert _deserializza_valore_csv('2026-09-26 09:20:56') == datetime(2026, 9, 26, 9, 20, 56)
    assert _deserializza_valore_csv('001_0-5 min') == '001_0-5 min'


def test_leggi_righe_export_csv_e_xlsx_stesso_risultato(tmp_path):
    intestazione = ['id_attivita', 'dist_m', 'timestamp']
    righe = [['att1', 12.5, datetime(2026, 1, 1, 12, 0, 0)]]

    percorso_csv = tmp_path / "a.csv"
    percorso_xlsx = tmp_path / "a.xlsx"
    _scrivi_csv(percorso_csv, intestazione, righe)
    _scrivi_xlsx(percorso_xlsx, intestazione, righe)

    int_csv, righe_csv = leggi_righe_export(percorso_csv)
    int_xlsx, righe_xlsx = leggi_righe_export(percorso_xlsx)

    assert int_csv == intestazione
    assert int_xlsx == intestazione
    assert righe_csv == righe
    assert righe_xlsx == righe


def test_leggi_righe_export_csv_vuoto_solleva_errore_chiaro(tmp_path):
    percorso = tmp_path / "vuoto.csv"
    percorso.write_text("")
    with pytest.raises(ValueError):
        leggi_righe_export(percorso)


def test_unisci_righe_export_unione_colonne_con_null():
    file_export = [
        (['id', 'a', 'b'], [['x1', 1, 2]]),
        (['id', 'a', 'c'], [['x2', 3, 4]]),
    ]
    intestazione, righe = unisci_righe_export(file_export)

    assert intestazione == ['id', 'a', 'b', 'c']
    assert righe == [
        ['x1', 1, 2, None],
        ['x2', 3, None, 4],
    ]


def test_unisci_output_csv(tmp_path, monkeypatch, capsys):
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(cartella_output / "att1.csv", ['id_attivita', 'speed_uff_kmh'], [['att1', 6.3]])
    _scrivi_xlsx(cartella_output / "att2.xlsx", ['id_attivita', 'speed_hav_kmh'], [['att2', 7.1]])

    monkeypatch.chdir(tmp_path)
    unisci_output('csv')
    out = capsys.readouterr().out

    # il merge va scritto in data/output/, la stessa cartella letta come
    # input (come --process-all, che ignora -o/--output-dir)
    file_merge = list(cartella_output.glob('merge_*.csv'))
    assert len(file_merge) == 1
    assert "Uniti 2 file" in out

    with open(file_merge[0], newline='', encoding='utf-8') as fh:
        righe = list(csv.reader(fh))
    # niente colonna 'timestamp' in questi file sintetici: ultima_attivita
    # resta vuota per tutti (non può stabilire un ordine temporale)
    assert righe[0] == ['id_attivita', 'speed_uff_kmh', 'speed_hav_kmh', 'ultima_attivita']
    assert righe[1] == ['att1', '6.3', '', '']
    assert righe[2] == ['att2', '', '7.1', '']


def test_unisci_output_nessun_file_non_crasha(tmp_path, monkeypatch, capsys):
    (tmp_path / "data" / "output").mkdir(parents=True)
    monkeypatch.chdir(tmp_path)
    unisci_output('csv')
    err = capsys.readouterr().err
    assert list(tmp_path.glob('data/output/merge_*.csv')) == []
    assert "Nessun file" in err


def test_unisci_output_file_non_valido_viene_saltato(tmp_path, monkeypatch, capsys):
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(cartella_output / "buono.csv", ['id_attivita'], [['att1']])
    (cartella_output / "rotto.csv").write_text("")

    monkeypatch.chdir(tmp_path)
    unisci_output('csv')
    captured = capsys.readouterr()

    assert "rotto.csv" in captured.err
    assert "Uniti 1 file" in captured.out


def test_unisci_output_esclude_un_merge_precedente_dall_input(tmp_path, monkeypatch, capsys):
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(cartella_output / "att1.csv", ['id_attivita'], [['att1']])

    monkeypatch.chdir(tmp_path)
    unisci_output('csv')  # primo merge: 1 file -> merge_1file_....csv
    capsys.readouterr()
    unisci_output('csv')  # secondo merge: deve ignorare il merge precedente
    out = capsys.readouterr().out

    assert "Uniti 1 file" in out  # non 2: il merge_1file_... non viene riletto


def test_cli_merge_output_end_to_end_xlsx(tmp_path, monkeypatch, capsys):
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(cartella_output / "att1.csv", ['id_attivita', 'speed_uff_kmh'], [['att1', 6.3]])
    _scrivi_csv(cartella_output / "att2.csv", ['id_attivita', 'speed_hav_kmh'], [['att2', 7.1]])

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr('sys.argv', ['milestats', '--merge-output', '-e', 'xlsx'])
    main()

    file_merge = list(cartella_output.glob('merge_*.xlsx'))
    assert len(file_merge) == 1
    ws = load_workbook(file_merge[0]).active
    assert [c.value for c in ws[1]] == [
        'id_attivita', 'speed_uff_kmh', 'speed_hav_kmh', 'ultima_attivita',
    ]


@pytest.mark.parametrize("argv_extra", [
    ['x.fit', '--merge-output'],
    ['--process-all', '--merge-output'],
    ['--merge-output', '--id', 'x'],
])
def test_merge_output_combinazioni_non_ammesse(argv_extra, monkeypatch):
    monkeypatch.setattr('sys.argv', ['milestats'] + argv_extra)
    with pytest.raises(SystemExit):
        main()


def test_merge_output_preserva_ts_min_attivita_per_attivita(tmp_path, monkeypatch):
    # dopo il merge, ts_min_attivita non deve mai mescolarsi tra attività
    # diverse: ogni riga mantiene il minimo della PROPRIA attività, anche
    # se accodata a righe di un'altra con un ts_min_attivita differente
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(
        cartella_output / "attA.csv",
        ['id_attivita', 'ts_min_attivita', 'timestamp'],
        [
            ['attA', '2026-01-01 08:00:00', '2026-01-01 08:00:00'],
            ['attA', '2026-01-01 08:00:00', '2026-01-01 08:05:00'],
        ],
    )
    _scrivi_csv(
        cartella_output / "attB.csv",
        ['id_attivita', 'ts_min_attivita', 'timestamp'],
        [['attB', '2026-02-02 09:00:00', '2026-02-02 09:00:00']],
    )

    monkeypatch.chdir(tmp_path)
    unisci_output('csv')

    file_merge = list(cartella_output.glob('merge_*.csv'))[0]
    with open(file_merge, newline='', encoding='utf-8') as fh:
        righe = list(csv.DictReader(fh))

    per_id = {riga['id_attivita']: riga['ts_min_attivita'] for riga in righe}
    assert per_id == {'attA': '2026-01-01 08:00:00', 'attB': '2026-02-02 09:00:00'}


def test_aggiungi_colonna_ultima_attivita_marca_solo_la_piu_recente():
    intestazione = ['id_attivita', 'timestamp']
    righe = [
        ['A', datetime(2026, 1, 1, 8, 0, 0)],
        ['A', datetime(2026, 1, 1, 8, 5, 0)],
        ['B', datetime(2026, 3, 3, 9, 0, 0)],   # attività più recente (inizia per ultima)
        ['C', datetime(2025, 12, 31, 7, 0, 0)],
    ]
    nuova_intestazione, nuove_righe = aggiungi_colonna_ultima_attivita(intestazione, righe)

    assert nuova_intestazione == ['id_attivita', 'timestamp', 'ultima_attivita']
    marcatori = {riga[0]: riga[-1] for riga in nuove_righe}
    # a parità di id_attivita il marcatore è coerente su tutte le righe
    assert [riga[-1] for riga in nuove_righe if riga[0] == 'A'] == [None, None]
    assert marcatori == {'A': None, 'B': 'X', 'C': None}


def test_aggiungi_colonna_ultima_attivita_nessuna_colonna_id_o_timestamp():
    intestazione = ['solo_una_colonna']
    righe = [['x'], ['y']]
    nuova_intestazione, nuove_righe = aggiungi_colonna_ultima_attivita(intestazione, righe)
    assert nuova_intestazione == ['solo_una_colonna', 'ultima_attivita']
    assert nuove_righe == [['x', None], ['y', None]]


def test_merge_output_aggiunge_colonna_ultima_attivita(tmp_path, monkeypatch):
    cartella_output = tmp_path / "data" / "output"
    cartella_output.mkdir(parents=True)
    _scrivi_csv(cartella_output / "A.csv", ['id_attivita', 'timestamp'],
                [['A', '2026-01-01 08:00:00']])
    _scrivi_csv(cartella_output / "B.csv", ['id_attivita', 'timestamp'],
                [['B', '2026-03-03 09:00:00'], ['B', '2026-03-03 09:05:00']])

    monkeypatch.chdir(tmp_path)
    unisci_output('csv')

    file_merge = list(cartella_output.glob('merge_*.csv'))[0]
    with open(file_merge, newline='', encoding='utf-8') as fh:
        righe = list(csv.DictReader(fh))

    assert [r['ultima_attivita'] for r in righe if r['id_attivita'] == 'A'] == ['']
    assert [r['ultima_attivita'] for r in righe if r['id_attivita'] == 'B'] == ['X', 'X']
